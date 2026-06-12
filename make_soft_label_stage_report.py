from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


PROJECT_DIR = Path(__file__).resolve().parent
REPORT_DIR = PROJECT_DIR / "stage_reports" / "soft_label_stage"
FIG_DIR = REPORT_DIR / "figures"
REPORT_PATH = REPORT_DIR / "soft_label_stage_report_2026-06-11.docx"


RESULTS = {
    "Baseline": {
        "accuracy": 94.24,
        "precision": 93.11,
        "recall": 95.05,
        "f1": 94.07,
        "false_alarm": 6.51,
        "miss": 4.95,
        "d_mae": 1.2704,
        "t_mae": 26.90,
        "cycle_acc": 96.24,
        "phase_mae": 0.0260,
        "high_conf_wrong": 69,
        "large_240": 62,
        "correct_cycle_t": 13.18,
        "wrong_cycle_t": 377.66,
    },
    "Soft 0.05": {
        "accuracy": 96.49,
        "precision": 95.23,
        "recall": 97.59,
        "f1": 96.40,
        "false_alarm": 4.53,
        "miss": 2.41,
        "d_mae": 1.2839,
        "t_mae": 25.94,
        "cycle_acc": 95.71,
        "phase_mae": 0.0271,
        "high_conf_wrong": 34,
        "large_240": 66,
        "correct_cycle_t": 12.84,
        "wrong_cycle_t": 318.15,
    },
    "Soft 0.10": {
        "accuracy": 97.32,
        "precision": 96.22,
        "recall": 98.28,
        "f1": 97.24,
        "false_alarm": 3.58,
        "miss": 1.72,
        "d_mae": 1.2810,
        "t_mae": 24.01,
        "cycle_acc": 95.21,
        "phase_mae": 0.0267,
        "high_conf_wrong": 13,
        "large_240": 56,
        "correct_cycle_t": 10.85,
        "wrong_cycle_t": 285.77,
    },
}


def find_font():
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/msyh.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return None


FONT_PATH = find_font()


def font(size, bold=False):
    if FONT_PATH:
        return ImageFont.truetype(FONT_PATH, size=size)
    return ImageFont.load_default()


def add_title(draw, text, width, y):
    draw.text((width // 2, y), text, anchor="ma", fill=(20, 20, 20), font=font(34))


def draw_axes(draw, x0, y0, x1, y1, ymax, ylabel):
    draw.line((x0, y0, x0, y1), fill=(40, 40, 40), width=2)
    draw.line((x0, y1, x1, y1), fill=(40, 40, 40), width=2)
    for i in range(6):
        v = ymax * i / 5
        y = y1 - (y1 - y0) * i / 5
        draw.line((x0 - 5, y, x1, y), fill=(225, 225, 225), width=1)
        draw.text((x0 - 12, y), f"{v:.0f}", anchor="rm", fill=(80, 80, 80), font=font(16))
    draw.text((x0, y0 - 34), ylabel, fill=(80, 80, 80), font=font(18))


def bar_chart(path, title, series, ymax, ylabel, value_fmt="{:.1f}", colors=None):
    width, height = 1300, 760
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    add_title(draw, title, width, 28)
    x0, y0, x1, y1 = 110, 110, width - 70, height - 135
    draw_axes(draw, x0, y0, x1, y1, ymax, ylabel)
    colors = colors or [(76, 120, 168), (245, 133, 24), (84, 162, 75)]
    groups = list(series.keys())
    labels = list(next(iter(series.values())).keys())
    group_width = (x1 - x0) / len(labels)
    bar_w = min(70, group_width / (len(groups) + 1.2))
    for li, label in enumerate(labels):
        center = x0 + group_width * (li + 0.5)
        for gi, group in enumerate(groups):
            value = series[group][label]
            h = (value / ymax) * (y1 - y0)
            bx0 = center - (len(groups) * bar_w) / 2 + gi * bar_w
            bx1 = bx0 + bar_w * 0.82
            by0 = y1 - h
            draw.rectangle((bx0, by0, bx1, y1), fill=colors[gi % len(colors)])
            draw.text(((bx0 + bx1) / 2, by0 - 8), value_fmt.format(value), anchor="mb", fill=(40, 40, 40), font=font(15))
        draw.text((center, y1 + 22), label, anchor="ma", fill=(40, 40, 40), font=font(17))
    lx, ly = x0, height - 75
    for gi, group in enumerate(groups):
        draw.rectangle((lx, ly, lx + 22, ly + 22), fill=colors[gi % len(colors)])
        draw.text((lx + 32, ly + 11), group, anchor="lm", fill=(40, 40, 40), font=font(18))
        lx += 220
    img.save(path)


def line_chart(path, title, data, ylabel, ymax=None):
    width, height = 1200, 720
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    add_title(draw, title, width, 28)
    x0, y0, x1, y1 = 120, 110, width - 90, height - 125
    ymax = ymax or max(max(vals) for vals in data.values()) * 1.15
    draw_axes(draw, x0, y0, x1, y1, ymax, ylabel)
    labels = list(RESULTS.keys())
    xs = [x0 + (x1 - x0) * i / (len(labels) - 1) for i in range(len(labels))]
    colors = [(76, 120, 168), (228, 87, 86), (84, 162, 75)]
    for si, (name, vals) in enumerate(data.items()):
        pts = []
        for x, v in zip(xs, vals):
            y = y1 - (v / ymax) * (y1 - y0)
            pts.append((x, y))
        draw.line(pts, fill=colors[si % len(colors)], width=4)
        for x, y, v in zip(xs, [p[1] for p in pts], vals):
            draw.ellipse((x - 6, y - 6, x + 6, y + 6), fill=colors[si % len(colors)])
            draw.text((x, y - 16), f"{v:.1f}", anchor="mb", fill=(40, 40, 40), font=font(15))
    for x, label in zip(xs, labels):
        draw.text((x, y1 + 25), label, anchor="ma", fill=(40, 40, 40), font=font(17))
    lx, ly = x0, height - 70
    for si, name in enumerate(data.keys()):
        draw.line((lx, ly + 11, lx + 35, ly + 11), fill=colors[si % len(colors)], width=4)
        draw.text((lx + 45, ly + 11), name, anchor="lm", fill=(40, 40, 40), font=font(18))
        lx += 320
    img.save(path)


def create_figures():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    classification = {
        name: {
            "Accuracy": vals["accuracy"],
            "Precision": vals["precision"],
            "Recall": vals["recall"],
            "F1": vals["f1"],
        }
        for name, vals in RESULTS.items()
    }
    bar_chart(FIG_DIR / "classification_comparison.png", "Classification Metrics Comparison", classification, 100, "Percent")

    regression = {
        name: {
            "D_MAE km": vals["d_mae"],
            "T_MAE/10": vals["t_mae"] / 10.0,
            "Cycle acc/100": vals["cycle_acc"] / 100.0,
        }
        for name, vals in RESULTS.items()
    }
    bar_chart(FIG_DIR / "regression_comparison.png", "Regression Metrics Comparison", regression, 3.2, "Scaled value", value_fmt="{:.2f}")

    tail = {
        name: {
            "High-conf wrong K": vals["high_conf_wrong"],
            ">240min errors": vals["large_240"],
        }
        for name, vals in RESULTS.items()
    }
    bar_chart(FIG_DIR / "tail_error_comparison.png", "Long-Tail Error Comparison", tail, 75, "Case count", value_fmt="{:.0f}")

    line_chart(
        FIG_DIR / "tmae_and_high_conf_trend.png",
        "T_MAE and High-Confidence Wrong-K Trend",
        {
            "T_MAE min": [RESULTS[k]["t_mae"] for k in RESULTS],
            "High-conf wrong K": [RESULTS[k]["high_conf_wrong"] for k in RESULTS],
        },
        "Value",
        ymax=75,
    )

    cycle_error = {
        "Baseline": {
            "K correct": RESULTS["Baseline"]["correct_cycle_t"],
            "K wrong": RESULTS["Baseline"]["wrong_cycle_t"],
        },
        "Soft 0.10": {
            "K correct": RESULTS["Soft 0.10"]["correct_cycle_t"],
            "K wrong": RESULTS["Soft 0.10"]["wrong_cycle_t"],
        },
    }
    bar_chart(FIG_DIR / "cycle_correct_wrong_tmae.png", "TCA Error by Cycle Prediction", cycle_error, 420, "T_MAE min")


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(9)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, item in enumerate(row):
            set_cell_text(cells[i], item)
    return table


def set_document_style(document):
    style = document.styles["Normal"]
    style.font.name = "SimSun"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    style.font.size = Pt(10.5)


def add_heading(document, text, level=1):
    p = document.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "SimHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    return p


def add_picture(document, image_path, caption):
    document.add_picture(str(image_path), width=Inches(6.2))
    p = document.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(90, 90, 90)


def create_report():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    create_figures()

    doc = Document()
    set_document_style(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("阶段性实验汇报：周期软标签对最近接近时间预测的改进")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "SimHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("基于轨道根数输入的大规模低轨空间目标碰撞风险快速预测模型").font.size = Pt(12)

    doc.add_paragraph("日期：2026-06-11")
    doc.add_paragraph("本阶段目标：在不改变输入维度、输出任务和 Transformer 主体结构的前提下，针对最近接近时间 Tc 的长尾误差问题，验证周期 K 相邻软标签是否能降低高置信周期错分并改善 T_MAE。")

    add_heading(doc, "1. 实验背景", 1)
    doc.add_paragraph("当前模型使用两颗目标的轨道六根数作为输入，并将角度项展开为 sin/cos，因此实际输入为 18 维。模型输出危险类别、最近接近距离 Dc，以及以平均轨道周期为参考的最近接近时间表示。")
    doc.add_paragraph("前期诊断发现，T_MAE 的主要来源不是所有样本的整体偏差，而是少数周期 K 预测错误样本造成的长尾误差。尤其是高置信错误周期样本中，大量错误集中在 K±1 的相邻周期。")

    add_heading(doc, "2. 方法改进", 1)
    doc.add_paragraph("原始周期分类使用 one-hot 硬标签，即真实周期 K 的标签为 1，其余周期为 0。这种做法会把 K±1 的相邻周期错误视为和远距离周期错误同等严重。")
    doc.add_paragraph("本阶段引入周期 K 相邻软标签：真实周期 K 保留主要权重，相邻周期 K-1 和 K+1 分配少量权重。这样模型仍然以真实周期为目标，但训练时对相邻周期保持一定物理连续性。")
    add_table(
        doc,
        ["版本", "周期标签设计", "含义"],
        [
            ["Baseline", "K=1.00，其余为0", "硬标签，严格分类"],
            ["Soft 0.05", "K=0.95，K±1总权重0.05", "轻微软化"],
            ["Soft 0.10", "K=0.90，K±1总权重0.10", "更明显地缓解相邻周期错分"],
        ],
    )

    add_heading(doc, "3. 关键结果", 1)
    add_table(
        doc,
        ["指标", "Baseline", "Soft 0.05", "Soft 0.10"],
        [
            ["Accuracy (%)", "94.24", "96.49", "97.32"],
            ["Precision (%)", "93.11", "95.23", "96.22"],
            ["Recall (%)", "95.05", "97.59", "98.28"],
            ["False alarm (%)", "6.51", "4.53", "3.58"],
            ["Miss rate (%)", "4.95", "2.41", "1.72"],
            ["D_MAE (km)", "1.2704", "1.2839", "1.2810"],
            ["T_MAE (min)", "26.90", "25.94", "24.01"],
            ["Cycle accuracy (%)", "96.24", "95.71", "95.21"],
            ["High-conf wrong K", "69", "34", "13"],
            [">240 min errors", "62", "66", "56"],
        ],
    )

    add_picture(doc, FIG_DIR / "classification_comparison.png", "图1 分类指标对比：软标签显著提高 Accuracy、Precision 和 Recall。")
    add_picture(doc, FIG_DIR / "regression_comparison.png", "图2 回归指标对比：Soft 0.10 获得最低 T_MAE，D_MAE 仅小幅变化。")
    add_picture(doc, FIG_DIR / "tail_error_comparison.png", "图3 长尾错误对比：Soft 0.10 将高置信错误周期从 69 降到 13。")
    add_picture(doc, FIG_DIR / "tmae_and_high_conf_trend.png", "图4 T_MAE 与高置信错误周期趋势：二者同步下降，说明长尾误差被有效压缩。")
    add_picture(doc, FIG_DIR / "cycle_correct_wrong_tmae.png", "图5 周期正确/错误时的时间误差：软标签降低了周期错误样本造成的平均时间误差。")

    add_heading(doc, "4. 结果分析", 1)
    doc.add_paragraph("Soft 0.10 是当前最优设置。相比 baseline，T_MAE 从 26.90 min 降至 24.01 min，高置信错误周期从 69 降至 13，漏报率从 4.95% 降至 1.72%。这说明周期软标签确实缓解了周期边界处的高置信错分。")
    doc.add_paragraph("Soft 0.05 虽然相对 baseline 有提升，但不如 Soft 0.10。它的周期准确率略高于 Soft 0.10，但高置信错误周期和 T_MAE 都更差，说明单纯追求严格周期准确率并不等价于更好的最近接近时间预测。")
    doc.add_paragraph("Soft 0.10 的周期准确率略低于 baseline，但这是可以接受的，因为模型对相邻周期不再过度自信，最终 T_MAE 和漏报率均明显改善。对本课题而言，最终风险筛查效果和 Tc 误差长尾更重要。")

    add_heading(doc, "5. 阶段性结论", 1)
    conclusions = [
        "当前阶段的主要瓶颈是 Tc 预测中的少数周期 K 错分样本，而不是所有样本的整体时间偏差。",
        "周期 K 相邻软标签是一种有效且可解释的改进方法，它符合轨道周期连续性的物理直觉。",
        "Soft 0.10 是当前推荐主线配置：真实周期权重 0.90，相邻周期总权重 0.10。",
        "该方法不改变输入维度、不改变输出任务、不改变 Transformer 主体，因此适合作为论文中的消融实验。"
    ]
    for item in conclusions:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "6. 后续建议", 1)
    suggestions = [
        "固定 Soft 0.10 作为当前主线 baseline，后续实验均与该版本对比。",
        "继续分析剩余 13 个高置信错误周期样本，判断是否存在特定轨道区域、Tc 区间或 Dc 区间集中现象。",
        "暂时不继续增加软标签强度到 0.15，避免周期头过软。",
        "后续若进一步优化，可考虑模型不确定性估计或专门针对高置信错误样本的诊断图。"
    ]
    for item in suggestions:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    create_report()
