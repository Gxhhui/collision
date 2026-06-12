import csv
from collections import Counter
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
CSV_PATH = ROOT / "logs" / "high_conf_wrong_cycle_cases.csv"
OUT_DIR = ROOT / "stage_reports" / "high_conf_wrong_cycle_analysis"
FIG_DIR = OUT_DIR / "figures"
DOCX_PATH = OUT_DIR / "high_conf_wrong_cycle_analysis_2026-06-11.docx"


NUMERIC_COLUMNS = [
    "index",
    "true_dc_km",
    "pred_dc_km",
    "true_tc_s",
    "pred_tc_s",
    "abs_time_error_min",
    "true_cycle",
    "raw_pred_cycle",
    "pred_cycle",
    "cycle_offset",
    "refined_applied",
    "raw_cycle_prob",
    "pred_cycle_prob",
    "true_cycle_prob",
    "top2_cycle_1",
    "top2_cycle_2",
    "true_phase",
    "pred_phase",
    "abs_phase_error",
    "refined_proxy_distance_km",
    "danger_prob",
]


def load_cases():
    if not CSV_PATH.exists():
        raise FileNotFoundError(f"Missing {CSV_PATH}. Run 02_train_transformer.py first.")
    rows = []
    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            parsed = {}
            for key, value in row.items():
                if key in NUMERIC_COLUMNS:
                    parsed[key] = float(value)
                else:
                    parsed[key] = value
            rows.append(parsed)
    if not rows:
        raise RuntimeError("No high-confidence wrong-cycle cases found.")
    return rows


def arr(rows, key):
    return np.array([float(r[key]) for r in rows], dtype=float)


def save_figures(rows):
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    plt.rcParams["figure.dpi"] = 160
    plt.rcParams["font.size"] = 10
    plt.rcParams["axes.unicode_minus"] = False

    abs_err = arr(rows, "abs_time_error_min")
    true_tc_h = arr(rows, "true_tc_s") / 3600.0
    true_dc = arr(rows, "true_dc_km")
    phase_err = arr(rows, "abs_phase_error")
    danger_prob = arr(rows, "danger_prob")
    raw_prob = arr(rows, "raw_cycle_prob")
    offsets = arr(rows, "cycle_offset").astype(int)

    order = np.argsort(-abs_err)
    labels = [str(int(rows[i]["index"])) for i in order]

    fig, ax = plt.subplots(figsize=(9, 4.8))
    ax.bar(np.arange(len(rows)), abs_err[order], color="#4c78a8")
    ax.set_xticks(np.arange(len(rows)))
    ax.set_xticklabels(labels, rotation=45, ha="right")
    ax.set_ylabel("Absolute TCA error (min)")
    ax.set_xlabel("Sample index")
    ax.set_title("High-Confidence Wrong-K Cases Ranked by TCA Error")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "ranked_time_error.png")
    plt.close(fig)

    counts = Counter(offsets)
    xs = sorted(counts)
    fig, ax = plt.subplots(figsize=(7, 4.6))
    ax.bar([str(x) for x in xs], [counts[x] for x in xs], color="#e45756")
    ax.set_xlabel("Predicted K - true K")
    ax.set_ylabel("Case count")
    ax.set_title("Cycle Offset Distribution")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "cycle_offset_distribution.png")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7, 4.8))
    sc = ax.scatter(true_tc_h, abs_err, c=np.abs(offsets), cmap="viridis", s=75, edgecolor="black", linewidth=0.4)
    ax.set_xlabel("True TCA (h)")
    ax.set_ylabel("Absolute TCA error (min)")
    ax.set_title("TCA Error vs True TCA")
    ax.grid(alpha=0.25)
    cbar = fig.colorbar(sc, ax=ax)
    cbar.set_label("|Cycle offset|")
    fig.tight_layout()
    fig.savefig(FIG_DIR / "error_vs_true_tca.png")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7, 4.8))
    sc = ax.scatter(true_dc, abs_err, c=phase_err, cmap="magma", s=75, edgecolor="black", linewidth=0.4)
    ax.set_xlabel("True Dc (km)")
    ax.set_ylabel("Absolute TCA error (min)")
    ax.set_title("TCA Error vs True Dc")
    ax.grid(alpha=0.25)
    cbar = fig.colorbar(sc, ax=ax)
    cbar.set_label("Absolute phase error")
    fig.tight_layout()
    fig.savefig(FIG_DIR / "error_vs_true_dc.png")
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(10, 4.3))
    axes[0].hist(raw_prob, bins=np.linspace(0.95, 1.0, 11), color="#54a24b", edgecolor="white")
    axes[0].set_xlabel("Raw cycle probability")
    axes[0].set_ylabel("Case count")
    axes[0].set_title("Cycle Confidence")
    axes[0].grid(axis="y", alpha=0.25)
    axes[1].hist(danger_prob, bins=np.linspace(0.0, 1.0, 11), color="#f58518", edgecolor="white")
    axes[1].set_xlabel("Danger probability")
    axes[1].set_title("Danger Classification Confidence")
    axes[1].grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "confidence_histograms.png")
    plt.close(fig)


def set_doc_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "SimSun"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    style.font.size = Pt(10.5)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "SimHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    return p


def add_picture(doc, name, caption):
    doc.add_picture(str(FIG_DIR / name), width=Inches(6.1))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(90, 90, 90)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(h)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(8.5)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
    return table


def create_doc(rows):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    save_figures(rows)

    abs_err = arr(rows, "abs_time_error_min")
    true_tc_h = arr(rows, "true_tc_s") / 3600.0
    true_dc = arr(rows, "true_dc_km")
    offsets = arr(rows, "cycle_offset").astype(int)
    phase_err = arr(rows, "abs_phase_error")
    danger_prob = arr(rows, "danger_prob")
    raw_prob = arr(rows, "raw_cycle_prob")

    doc = Document()
    set_doc_style(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("高置信错误周期样本诊断报告")
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "SimHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Soft 0.10 周期相邻软标签主线模型").font.size = Pt(12)

    doc.add_paragraph("日期：2026-06-11")
    doc.add_paragraph("分析对象：logs/high_conf_wrong_cycle_cases.csv 中剩余高置信周期错分样本。当前主线模型已经将高置信错误周期样本从 baseline 的 69 个降低到 13 个，本报告分析这 13 个样本的剩余错误特征。")

    heading(doc, "1. 总体统计", 1)
    add_table(
        doc,
        ["项目", "数值"],
        [
            ["样本数", len(rows)],
            ["平均时间误差 min", f"{np.mean(abs_err):.2f}"],
            ["中位时间误差 min", f"{np.median(abs_err):.2f}"],
            ["最大时间误差 min", f"{np.max(abs_err):.2f}"],
            ["平均 true Dc km", f"{np.mean(true_dc):.2f}"],
            ["中位 true Dc km", f"{np.median(true_dc):.2f}"],
            ["平均 true TCA h", f"{np.mean(true_tc_h):.2f}"],
            ["中位 true TCA h", f"{np.median(true_tc_h):.2f}"],
            ["平均 phase error", f"{np.mean(phase_err):.3f}"],
            ["平均 raw cycle prob", f"{np.mean(raw_prob):.3f}"],
            ["平均 danger prob", f"{np.mean(danger_prob):.3f}"],
        ],
    )

    heading(doc, "2. 错误样本排序", 1)
    top_rows = sorted(rows, key=lambda r: r["abs_time_error_min"], reverse=True)
    add_table(
        doc,
        ["index", "Dc km", "TCA err min", "true K", "pred K", "offset", "phase err", "cycle prob", "danger prob"],
        [
            [
                int(r["index"]),
                f"{r['true_dc_km']:.2f}",
                f"{r['abs_time_error_min']:.1f}",
                int(r["true_cycle"]),
                int(r["pred_cycle"]),
                int(r["cycle_offset"]),
                f"{r['abs_phase_error']:.3f}",
                f"{r['raw_cycle_prob']:.3f}",
                f"{r['danger_prob']:.3f}",
            ]
            for r in top_rows
        ],
    )
    add_picture(doc, "ranked_time_error.png", "图1 高置信错误周期样本按时间误差排序。")

    heading(doc, "3. 周期错分模式", 1)
    offset_counter = Counter(offsets)
    offset_text = "，".join([f"offset={k}: {v}" for k, v in sorted(offset_counter.items())])
    doc.add_paragraph(f"周期错分分布为：{offset_text}。其中相邻周期错误仍占主要部分，但已经从 baseline 中的大量 K±1 错分降低到较小规模。")
    add_picture(doc, "cycle_offset_distribution.png", "图2 周期偏移分布。")

    heading(doc, "4. 与真实时间和距离的关系", 1)
    doc.add_paragraph("剩余错误样本的真实 TCA 中位值约为 23 小时，说明错误并不只出现在 48 小时末端；同时 Dc 中位值约为 4 km，说明部分非常近距离样本仍可能出现周期错分。")
    add_picture(doc, "error_vs_true_tca.png", "图3 时间误差与真实 TCA 的关系。")
    add_picture(doc, "error_vs_true_dc.png", "图4 时间误差与真实 Dc 的关系。")

    heading(doc, "5. 置信度特征", 1)
    doc.add_paragraph("这些样本的 raw cycle probability 均大于 0.95，说明它们不是普通不确定样本，而是模型非常自信但周期判断错误的样本。部分样本的 danger probability 较低，说明分类头和周期头之间仍存在不一致。")
    add_picture(doc, "confidence_histograms.png", "图5 周期置信度与危险分类置信度分布。")

    heading(doc, "6. 结论与下一步建议", 1)
    conclusions = [
        "Soft 0.10 已经显著压低高置信周期错分，但剩余 13 个样本仍然贡献较大的 T_MAE 长尾。",
        "剩余错误中仍包含 K±1 相邻周期错误，也包含少数更远的周期偏移，说明单纯软标签已经接近当前方案的收益边界。",
        "部分样本 danger probability 不高但周期头极高置信，后续可考虑分类置信度与周期置信度联合诊断。",
        "下一步不建议继续盲目增大软标签权重，建议固定 Soft 0.10 作为主线，进一步做不确定性估计或错误样本专门分析。",
    ]
    for item in conclusions:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


def main():
    rows = load_cases()
    create_doc(rows)


if __name__ == "__main__":
    main()
